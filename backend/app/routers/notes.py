import io
from datetime import datetime
from typing import List

from fastapi import APIRouter, Depends, HTTPException, UploadFile, File, Form
from sqlalchemy.orm import Session

from ..database import get_db
from ..models import Note, User
from ..schemas import NoteCreate, NoteResponse
from ..dependencies import get_current_user

router = APIRouter(tags=["Notes"])


def extract_text_from_pdf(data: bytes) -> str:
    try:
        import pypdf
        reader = pypdf.PdfReader(io.BytesIO(data))
        text = "\n".join(p.extract_text() or "" for p in reader.pages)
        if not text.strip():
            raise HTTPException(400, "PDF appears scanned — no text extractable")
        return text
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(400, f"Could not read PDF: {e}")


def extract_text_from_ppt(data: bytes) -> str:
    try:
        from pptx import Presentation
        prs = Presentation(io.BytesIO(data))
        parts = [shape.text.strip() for slide in prs.slides for shape in slide.shapes if hasattr(shape, "text") and shape.text.strip()]
        text = "\n".join(parts)
        if not text.strip():
            raise HTTPException(400, "PowerPoint has no readable text")
        return text
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(400, f"Could not read PowerPoint: {e}")


def make_note_content(raw_text: str, title: str) -> str:
    """Store the extracted document text as the note content (AI summarisation removed)."""
    text = raw_text.strip()
    if not text:
        return f"(No text could be extracted from '{title}')"
    # Cap stored content to keep notes manageable
    return text[:5000]


@router.get("/notes", response_model=List[NoteResponse])
def get_notes(db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    return db.query(Note).filter(Note.user_id == current_user.id).order_by(Note.id.desc()).all()


@router.post("/notes", response_model=NoteResponse, status_code=201)
def create_note(payload: NoteCreate, db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    note = Note(title=payload.title, content=payload.content, subject=payload.subject,
                file_type="manual", created_at=str(datetime.now().date()), user_id=current_user.id)
    db.add(note); db.commit(); db.refresh(note)
    return note


@router.post("/notes/upload", response_model=NoteResponse, status_code=201)
async def upload_note(file: UploadFile = File(...), subject: str = Form(default=""),
                      db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    filename = file.filename or ""
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if ext not in ("pdf", "ppt", "pptx"):
        raise HTTPException(400, "Only PDF or PowerPoint files supported")
    data = await file.read()
    if len(data) > 20 * 1024 * 1024:
        raise HTTPException(400, "File too large — max 20MB")
    raw_text = extract_text_from_pdf(data) if ext == "pdf" else extract_text_from_ppt(data)
    title = filename.rsplit(".", 1)[0].replace("_", " ").replace("-", " ").title()
    summary = make_note_content(raw_text, title)
    note = Note(title=title, content=summary, subject=subject or None,
                file_name=filename, file_type="pdf" if ext == "pdf" else ext,
                raw_text=raw_text[:20000], created_at=str(datetime.now().date()), user_id=current_user.id)
    db.add(note); db.commit(); db.refresh(note)
    return note


@router.put("/notes/{note_id}", response_model=NoteResponse)
def update_note(note_id: int, payload: NoteCreate, db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    note = db.query(Note).filter(Note.id == note_id, Note.user_id == current_user.id).first()
    if not note: raise HTTPException(404, "Note not found")
    note.title = payload.title; note.content = payload.content; note.subject = payload.subject
    db.commit(); db.refresh(note)
    return note


@router.delete("/notes/{note_id}")
def delete_note(note_id: int, db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    note = db.query(Note).filter(Note.id == note_id, Note.user_id == current_user.id).first()
    if not note: raise HTTPException(404, "Note not found")
    db.delete(note); db.commit()
    return {"message": "Deleted"}
